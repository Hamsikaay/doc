# tasks.py
"""
Celery tasks:
- process_file: extract text from uploaded bytes, store Document and Chunks in DB,
                then enqueue ingest_document to create embeddings and insert into FAISS.
- ingest_document: read chunks from DB, embed them, and add to vectorstore.
"""
import io
import uuid
import time
from rag.celery_app import celery
from rag.db import SessionLocal, Document, Chunk
from rag.chunker import chunk_text
from rag.vectorstore import vstore
from rag.embeddings import embed_text
import pdfplumber

@celery.task(bind=True)
def process_file(self, content_bytes, filename: str = None, user_id: int = None):
    """
    1. Extract text (PDF-aware)
    2. Persist Document row with user_id
    3. Chunk text and persist Chunk rows
    4. Enqueue ingest_document for embeddings
    Returns doc_id and ingestion task id (if created)
    """
    # 1) extract text
    text = ""
    try:
        if filename and filename.lower().endswith(".pdf"):
            # PDF extraction
            with pdfplumber.open(io.BytesIO(content_bytes)) as pdf:
                pages = []
                for p in pdf.pages:
                    pages.append(p.extract_text() or "")
                text = "\n".join(pages)
        elif filename and filename.lower().endswith((".docx", ".doc")):
            # DOCX extraction
            try:
                from docx import Document as DocxDocument
                doc = DocxDocument(io.BytesIO(content_bytes))
                paragraphs = [para.text for para in doc.paragraphs]
                text = "\n".join(paragraphs)
            except ImportError:
                print("WARNING: python-docx not installed. Install with: pip install python-docx")
                text = ""
        else:
            # Plain text
            text = content_bytes.decode("utf-8")
    except Exception as e:
        print(f"ERROR extracting text from {filename}: {e}")
        # Try fallback to UTF-8
        try:
            text = content_bytes.decode("utf-8")
        except:
            text = ""
    
    print(f"Extracted {len(text)} characters from {filename}")

    # 2) save Document
    doc_id = str(uuid.uuid4())
    from datetime import datetime
    db = SessionLocal()
    doc = Document(
        doc_id=doc_id, 
        filename=filename or "", 
        text=text,
        user_id=user_id,
        created_at=datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    )
    db.add(doc)
    db.commit()

    # 3) chunk and save chunks
    num_chunks = 0
    for i, ch in enumerate(chunk_text(text or "", chunk_size=1000, overlap=200)):
        c = Chunk(doc_id=doc_id, chunk_index=i, text=ch)
        db.add(c)
        num_chunks = i + 1
    db.commit()
    db.close()

    # 4) enqueue ingestion task to build embeddings + vector index
    ingest_task = ingest_document.apply_async(args=[doc_id])

    return {"doc_id": doc_id, "chunks": num_chunks, "ingest_task_id": ingest_task.id}

@celery.task(bind=True)
def ingest_document(self, doc_id: str):
    """
    Read chunks for doc_id, embed each chunk, add to FAISS vectorstore.
    Returns number of added vectors.
    """
    db = SessionLocal()
    chunks = db.query(Chunk).filter(Chunk.doc_id == doc_id).order_by(Chunk.chunk_index).all()
    added = 0
    for c in chunks:
        vec = embed_text(c.text)
        meta = {"doc_id": doc_id, "chunk_index": c.chunk_index, "text_preview": c.text[:300]}
        vstore.add(vec, meta)
        added += 1
    db.close()
    return {"doc_id": doc_id, "added": added}
